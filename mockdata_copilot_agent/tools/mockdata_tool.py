"""
mockdata_tool — pure-function BigQuery GoogleSQL mock-data generator
====================================================================

This module generates **BigQuery GoogleSQL** INSERT text and the read-only
existence-lookup SELECT for the L'Oreal CDS B2C DV (SEA) replica dataset. It is a
PURE-function module: NO LLM calls, NO side effects beyond reading the committed
column reference (`.docs/dv_columns.json`). All functions RETURN strings; nothing
is ever printed to stdout (MCP-stdio safety). Any diagnostics go to stderr.

The agent NEVER executes the generated INSERTs — it presents them for a human to
run against BigQuery. The only SQL that touches a live database is the read-only
existence SELECT built by `build_lookup_sql`, which is blocklist-clean by
construction.

Confirmed decisions (locked by user 2026-06-24):
  - A1 [SQL dialect = GoogleSQL — CONFIRMED 2026-06-24]: Generated SQL is BigQuery
    GoogleSQL targeting the DV dataset
    `oa-apmena-spacecds-ap-dv.space_cds_b2c_ds_replica_ase1_dv_public`.
    Fully-qualified backtick table refs; STRING literals single-quoted; TIMESTAMP
    columns rendered as TIMESTAMP 'YYYY-MM-DD HH:MM:SS+00'; INT64/NUMERIC/FLOAT64
    unquoted; BOOL as true/false; bare NULL for empty cells. NO Postgres
    ::uuid / ::timestamptz casts (the .docs sample is Postgres and is the
    column/value reference ONLY, not the dialect).
  - A2 / A3 [column types from committed dv_columns.json — live-fetched]: value
    rendering is TYPE-DRIVEN from `.docs/dv_columns.json` (STRING vs INT64 vs
    TIMESTAMP), not name heuristics.
  - A4 [existence key = lucid (+ optional consumerkey/signature/isocountry
    disambiguators)]: lucid is not unique across brands.
  - A5 [UUID family]: synthetic UUID family
    `1b48bdbb-bdbb-7b6f-GGGG-NNNNNNNNNNNN` / `6d74baaa-baaa-7a19-899c-NNNNNNNNNNNN`;
    the trailing 12-digit counter is the uniqueness mechanism for extsystemidskey.
  - A8 [provenance]: MOCK_DATA / MANUAL markers preserved verbatim in every row.
  - A9 [synthetic PII]: consumer_pii PII is obviously fake
    (mock+<id>@example.com, fake names).
"""

from __future__ import annotations

import copy
import json
import re
import sys
from pathlib import Path

# ── DV dataset identifiers ──────────────────────────────────────────────────
# These are SQL identifiers embedded in the generated / lookup TEXT, not a
# configuration of the BigQuery client. The executor's billing project still
# comes from config (BQ_PROJECT); see 02-RESEARCH Q1. Optional kwargs let
# callers override.
_DV_PROJECT = "oa-apmena-spacecds-ap-dv"
_DV_DATASET = "space_cds_b2c_ds_replica_ase1_dv_public"

# Legacy SAMPLE column list (21 cols, from the .docs xlsx/.sql mock-up). This is a
# FALLBACK ONLY. The authoritative external_system_ids column set is fetched live
# into .docs/dv_columns.json: the live BigQuery DV table has 18 columns, and the 3
# trailing sample columns (recordstatus, country_iso_3166_1, consumersourcekey) do
# NOT exist on it. Generation is schema-driven from dv_columns.json; this tuple is
# used only when that reference is missing.
EXTERNAL_SYSTEM_IDS_COLUMNS = (
    "extsystemidskey",
    "consumerkey",
    "isocountry",
    "signature",
    "lucid",
    "sourcesystemid",
    "sourceconsumerid",
    "creationlocationid",
    "locationid",
    "creationsalesrepid",
    "salesrepid",
    "createddatetime",
    "status",
    "lastlogindatetime",
    "createdsourcesystemid",
    "lastupdatedsourcesystemid",
    "syscreateddatetime",
    "sysupdatedatetime",
    "recordstatus",
    "country_iso_3166_1",
    "consumersourcekey",
)

# Path to the committed column reference, resolved relative to this module's
# package (robust to CWD).
_DV_COLUMNS_JSON = Path(__file__).resolve().parent.parent / ".docs" / "dv_columns.json"

_NUMERIC_TYPES = {"INT64", "INTEGER", "INT", "NUMERIC", "BIGNUMERIC", "FLOAT64", "FLOAT", "DECIMAL"}
_BOOL_TYPES = {"BOOL", "BOOLEAN"}
# Complex types have no safe scalar literal — rendered as NULL. The live DV
# consumer_pii has a datastream_metadata STRUCT<uuid STRING, source_timestamp INT64>.
_COMPLEX_TYPES = {"STRUCT", "ARRAY", "JSON", "GEOGRAPHY", "RECORD"}


def _stderr(msg: str) -> None:
    """Emit a diagnostic to stderr only (never stdout — MCP-stdio safety)."""
    print(msg, file=sys.stderr)


def _table_ref(table: str, project: str = _DV_PROJECT, dataset: str = _DV_DATASET) -> str:
    """Return a fully-qualified backtick GoogleSQL table ref."""
    return f"`{project}.{dataset}.{table}`"


def load_column_types(table: str) -> dict:
    """Pure read of .docs/dv_columns.json.

    Returns {column_name: DATA_TYPE_UPPER} for the requested table. Empty dict if
    the file or table entry is missing (callers degrade to STRING). NO live query.
    """
    try:
        data = json.loads(_DV_COLUMNS_JSON.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        _stderr(f"[mockdata_tool] could not read {_DV_COLUMNS_JSON}: {exc}")
        return {}
    cols = data.get("tables", {}).get(table, [])
    return {c["name"]: str(c.get("data_type", "")).upper() for c in cols}


def _column_order(table: str, fallback: tuple) -> list:
    """Return the ordinal column order for a table from dv_columns.json,
    falling back to the provided tuple if the reference is missing."""
    try:
        data = json.loads(_DV_COLUMNS_JSON.read_text(encoding="utf-8"))
        cols = data.get("tables", {}).get(table, [])
        if cols:
            return [c["name"] for c in cols]
    except (OSError, ValueError):
        pass
    return list(fallback)


# ── ISO timestamp normalization ─────────────────────────────────────────────

def _normalize_timestamp(value: str) -> str:
    """Normalize an ISO-Z timestamp string to 'YYYY-MM-DD HH:MM:SS[.ffffff]+00'.

    Examples:
        '2026-06-19T00:00:00Z'            -> '2026-06-19 00:00:00+00'
        '2026-06-23T09:48:27.302993Z'     -> '2026-06-23 09:48:27.302993+00'
        '1900-01-01 00:00:00+00'          -> '1900-01-01 00:00:00+00' (idempotent)
    """
    s = str(value).strip()
    # Replace the date/time 'T' separator with a space.
    s = s.replace("T", " ")
    # Normalize the trailing zone marker to +00.
    if s.endswith("Z") or s.endswith("z"):
        s = s[:-1] + "+00"
    elif s.endswith("+00:00"):
        s = s[:-6] + "+00"
    elif not re.search(r"[+-]\d{2}(:?\d{2})?$", s):
        # No zone offset present at all — append +00.
        s = s + "+00"
    return s


def _normalize_datetime(value: str) -> str:
    """Normalize to 'YYYY-MM-DD HH:MM:SS[.ffffff]' with no zone offset (DATETIME)."""
    s = _normalize_timestamp(value)
    s = re.sub(r"[+-]\d{2}(:?\d{2})?$", "", s).strip()
    return s


def _normalize_date(value: str) -> str:
    """Normalize to 'YYYY-MM-DD' (DATE)."""
    s = str(value).strip().replace("T", " ")
    return s.split(" ")[0].split("Z")[0]


def _sq(value) -> str:
    """Single-quote a literal, doubling embedded single quotes."""
    return "'" + str(value).replace("'", "''") + "'"


def _render_value(value, data_type: str) -> str:
    """Type-driven GoogleSQL literal rendering.

    - None or "" -> bare NULL
    - STRING / BYTES / unknown -> single-quoted, embedded ' doubled
    - TIMESTAMP -> TIMESTAMP 'YYYY-MM-DD HH:MM:SS[.ffffff]+00'
    - DATE -> DATE 'YYYY-MM-DD' ; DATETIME -> DATETIME '<norm no offset>'
    - INT64/NUMERIC/FLOAT64/... -> unquoted numeric form (strips surrounding
      quotes if the seed gave a numeric string)
    - BOOL -> true / false
    """
    if value is None or value == "":
        return "NULL"

    # Normalize to the base type token: 'STRING(50)' -> 'STRING',
    # 'STRUCT<...>' -> 'STRUCT', 'ARRAY<...>' -> 'ARRAY'. The live DV schema uses
    # parameterized STRING(N) and a datastream_metadata STRUCT.
    base = (data_type or "").upper().split("(")[0].split("<")[0].strip()

    if base in _COMPLEX_TYPES:
        # No safe scalar literal for STRUCT/ARRAY/JSON/GEOGRAPHY -> NULL.
        return "NULL"
    if base == "TIMESTAMP":
        return f"TIMESTAMP {_sq(_normalize_timestamp(value))}"
    if base == "DATE":
        return f"DATE {_sq(_normalize_date(value))}"
    if base == "DATETIME":
        return f"DATETIME {_sq(_normalize_datetime(value))}"
    if base in _BOOL_TYPES:
        if isinstance(value, bool):
            return "true" if value else "false"
        return "true" if str(value).strip().lower() in ("true", "1", "t", "yes") else "false"
    if base in _NUMERIC_TYPES:
        # Strip surrounding quotes / whitespace; emit the bare number.
        return str(value).strip().strip("'").strip('"')

    # STRING, BYTES, or unknown -> quoted literal.
    return _sq(value)


# ── extsystemidskey counter increment (A5) ──────────────────────────────────

def _bump_extsystemidskey(key: str, offset: int) -> str:
    """Increment the trailing 12-digit counter of an extsystemidskey by `offset`.

    Pattern: '1b48bdbb-bdbb-7b6f-GGGG-NNNNNNNNNNNN'. The final 12-digit group is
    the per-row uniqueness counter. If the key does not match the expected shape,
    append a deterministic '-<offset>' suffix so the PK is still unique.
    """
    m = re.match(r"^(.*-)(\d{12})$", str(key))
    if not m:
        if offset == 0:
            return str(key)
        return f"{key}-{offset}"
    prefix, counter = m.group(1), int(m.group(2))
    return f"{prefix}{counter + offset:012d}"


def _vary_value(strategy, row_index: int, seed_value):
    """Resolve a `vary` strategy for a given row index.

    - list  -> cycle by row index
    - scalar (str/int) -> used verbatim for every row (deterministic constant)
    """
    if isinstance(strategy, (list, tuple)):
        if not strategy:
            return seed_value
        return strategy[row_index % len(strategy)]
    return strategy


def _render_row(row: dict, columns, column_types: dict) -> str:
    """Render one row dict to a parenthesised GoogleSQL VALUES tuple."""
    rendered = [
        _render_value(row.get(col), column_types.get(col, ""))
        for col in columns
    ]
    return "(" + ", ".join(rendered) + ")"


def _build_insert(table: str, columns, rows: list, column_types: dict,
                  project: str = _DV_PROJECT, dataset: str = _DV_DATASET) -> str:
    """Assemble a full GoogleSQL INSERT block matching the .docs layout."""
    cols_block = ",\n  ".join(columns)
    value_tuples = ",\n  ".join(_render_row(r, columns, column_types) for r in rows)
    return (
        f"INSERT INTO {_table_ref(table, project, dataset)} (\n"
        f"  {cols_block}\n"
        f") VALUES\n"
        f"  {value_tuples};"
    )


# Seed columns that appear in the .docs sample/xlsx but are NOT in the live
# BigQuery external_system_ids table. Per user decision 2026-06-24 ("Match BQ +
# warn"), these are dropped from generated INSERTs and explained in a comment.
_SAMPLE_ONLY_NOTES = {
    "recordstatus": "not in BigQuery external_system_ids (exists on consumer_pii)",
    "country_iso_3166_1": "not in BigQuery external_system_ids (exists on consumer_pii)",
    "consumersourcekey": "not in BigQuery external_system_ids or consumer_pii",
}


def _dropped_columns_comment(dropped: list) -> str:
    """Return a leading GoogleSQL comment block listing seed columns dropped
    because they are not in the live external_system_ids table (empty string if
    none). User decision 2026-06-24: Match BQ + warn on extras."""
    if not dropped:
        return ""
    lines = ["-- Dropped seed column(s) not present in BigQuery external_system_ids:"]
    for col in dropped:
        note = _SAMPLE_ONLY_NOTES.get(col, "not in BigQuery external_system_ids")
        lines.append(f"--   {col} — {note}")
    return "\n".join(lines) + "\n"


def generate_external_system_ids(seed: dict, vary: dict, count: int,
                                 column_types: dict | None = None,
                                 columns: list | None = None,
                                 project: str = _DV_PROJECT,
                                 dataset: str = _DV_DATASET) -> dict:
    """Pure. Generate `count` external_system_ids rows from a seed + variation spec.

    Returns {"sql": str, "rows": list[dict], "lucids": list[str],
             "consumerkeys": list[str], "dropped": list[str]}.

    Schema-driven: the column set defaults to the live BigQuery DV schema (18 cols)
    from .docs/dv_columns.json. Seed keys not in that schema (e.g. the sample-only
    recordstatus / country_iso_3166_1 / consumersourcekey) are DROPPED from the
    INSERT and explained in a leading comment (user decision 2026-06-24: Match BQ +
    warn). `vary` maps column_name -> value-list (cycled by row index) or scalar
    (constant). `extsystemidskey` is ALWAYS varied (trailing 12-digit counter ++)
    to keep the PK unique (A5). `column_types`/`columns` may be injected for tests.
    NO LLM, NO I/O beyond the columns reference.
    """
    column_types = column_types or load_column_types("external_system_ids")
    vary = vary or {}
    columns = columns or _column_order("external_system_ids", EXTERNAL_SYSTEM_IDS_COLUMNS)

    seed_key = seed.get("extsystemidskey")
    rows: list[dict] = []
    for i in range(count):
        row = copy.deepcopy(seed)
        # extsystemidskey ALWAYS varied (A5) — increment trailing counter per row.
        if seed_key is not None:
            row["extsystemidskey"] = _bump_extsystemidskey(seed_key, i)
        # Apply explicit variation strategies.
        for col, strategy in vary.items():
            if col == "extsystemidskey":
                continue  # PK is owned by the counter logic above.
            row[col] = _vary_value(strategy, i, seed.get(col))
        rows.append(row)

    # Match BQ + warn: drop seed keys absent from the live schema, note them.
    dropped = [k for k in seed.keys() if k not in columns]
    sql_body = _build_insert("external_system_ids", columns, rows, column_types,
                             project, dataset)
    sql = _dropped_columns_comment(dropped) + sql_body

    lucids = list(dict.fromkeys(r.get("lucid") for r in rows if r.get("lucid")))
    consumerkeys = list(dict.fromkeys(r.get("consumerkey") for r in rows if r.get("consumerkey")))
    return {"sql": sql, "rows": rows, "lucids": lucids, "consumerkeys": consumerkeys,
            "dropped": dropped}


def generate_consumer_pii(consumers: list, column_types: dict | None = None,
                          project: str = _DV_PROJECT,
                          dataset: str = _DV_DATASET) -> dict:
    """Pure. Generate synthetic consumer_pii INSERT rows for ABSENT consumers.

    `consumers` = list of {"lucid","consumerkey","signature","isocountry", ...}
    found absent in the existence lookup. Returns {"sql": str, "rows": list[dict]}.
    Synthetic fake PII only (mock+<id>@example.com, fake names, MOCK_DATA markers)
    across the committed DV consumer_pii column list (A9). `column_types` defaults
    to load_column_types('consumer_pii').
    """
    column_types = column_types or load_column_types("consumer_pii")
    columns = _column_order("consumer_pii", ())
    if not columns:
        # Degrade to the keys present on the consumers + known fake fields.
        columns = ["consumerkey", "isocountry", "signature", "lucid", "status",
                   "firstname", "lastname", "email", "recordstatus"]

    rows: list[dict] = []
    for idx, c in enumerate(consumers):
        ident = c.get("lucid") or c.get("consumerkey") or str(idx)
        row = {col: None for col in columns}
        # Carry over the known identifying keys verbatim.
        for k in ("lucid", "consumerkey", "signature", "isocountry"):
            if k in row and c.get(k) is not None:
                row[k] = c.get(k)
        # Obviously-fake synthetic PII (A9).
        if "firstname" in row:
            row["firstname"] = "Mock"
        if "lastname" in row:
            row["lastname"] = "Tester"
        if "email" in row:
            row["email"] = f"mock+{ident}@example.com"
        if "mobile" in row:
            row["mobile"] = "+0000000000"
        if "phone" in row:
            row["phone"] = "+0000000000"
        # Provenance / status markers (A8).
        if "status" in row and row["status"] is None:
            row["status"] = "Active"
        if "consumertype" in row:
            row["consumertype"] = "MOCK_DATA"
        if "recordstatus" in row:
            row["recordstatus"] = 0
        rows.append(row)

    sql = _build_insert("consumer_pii", columns, rows, column_types, project, dataset)
    return {"sql": sql, "rows": rows}


def build_lookup_sql(lucid: str, consumerkey: str | None = None,
                     signature: str | None = None, isocountry: str | None = None,
                     project: str = _DV_PROJECT, dataset: str = _DV_DATASET) -> str:
    """Pure. Return the read-only DV consumer_pii existence SELECT (GoogleSQL).

    Keys on `lucid` (A4) with optional AND clauses for provided disambiguators.
    Targets the fully-qualified consumer_pii ref so _pick_executor selects the DV
    executor. All literals single-quoted. Passes SafeBigQueryExecutor.validate_query.
    """
    where = [f"lucid = {_sq(lucid)}"]
    if consumerkey:
        where.append(f"consumerkey = {_sq(consumerkey)}")
    if signature:
        where.append(f"signature = {_sq(signature)}")
    if isocountry:
        where.append(f"isocountry = {_sq(isocountry)}")
    where_clause = "\n  AND ".join(where)
    return (
        "SELECT consumerkey, lucid, signature, isocountry, status\n"
        f"FROM {_table_ref('consumer_pii', project, dataset)}\n"
        f"WHERE {where_clause}\n"
        "LIMIT 5"
    )


def parse_seed_xlsx(path: str) -> dict:
    """Optional convenience: read a seed row from the .docs xlsx layout.

    Header row 0 = column names, row 1 = seed values. Paste-JSON is the primary
    input (A7); openpyxl is imported lazily so a missing dependency never blocks
    the core flow.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - optional path
        raise RuntimeError(
            "openpyxl is required for parse_seed_xlsx; paste the seed as JSON instead."
        ) from exc
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=1, max_row=2, values_only=True))
    headers = [str(h).strip() for h in rows[0] if h is not None]
    values = rows[1] if len(rows) > 1 else ()
    return {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}
